#!/usr/bin/env python3
"""
Build an import file for the Trifecta tracker from the legacy sources:

  * PROJECT TRACKER.xlsx, "2026" tab (one row per project, carry-over 25xxx projects included)
  * a TSV export of Google Drive project folders (Folder Name / Location / URL), matched by project #
  * signed OpenSolar contract PDFs (optional), matched to projects by customer name

Output:
  * import.json  - consumed by `php bin/console import:projects import.json`
  * report.md    - everything a person should review: guesses, mismatches, skipped rows

Nothing here talks to the database. Customer data stays in the output files; do not commit them.

Usage:
  python3 tools/import/build_import.py --xlsx "PROJECT TRACKER.xlsx" --drive folders.tsv \
	  --contracts contracts/ --out import/import.json --report import/report.md
Requires: openpyxl, and pdftotext (poppler-utils) for contracts.

Record-only import of the older "2024 & 2025" tab (history for reports; dates only, no payments,
existing project numbers skipped, nothing existing changed):
  python3 tools/import/build_import.py --older --xlsx "PROJECT TRACKER.xlsx" --contracts contracts/ \
	  --db trifecta-backup.sqlite --out import/older.json --report import/older-report.md
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import re
import subprocess
import sys
from pathlib import Path

import openpyxl

TODAY = dt.date.today()

# Per-project data lives in import/overrides.json, which is gitignored: customer names, site
# addresses, hand-read contracts (contact, price), municipality checks and Greg's answers to
# review notes. Customer data never goes in the repo. Without the file the builder still runs,
# just with no per-project fixes. Tables (project numbers as text):
#   SKIP_AS_SERVICE       rows that are really service work
#   CUSTOMER_ALIASES      [sheet-name regex, customer name] for projects sharing a customer
#   CENSUS_MUNI           project -> [municipality, county] from the Census geocoder
#   SITE_OVERRIDE         project -> site address Greg supplied
#   CONFIRMED_MUNI        project -> [municipality, county] Greg confirmed
#   DATE_FIX              "project|task path" -> corrected date
#   MANUAL_PROJECTS       projects with a Drive folder but no tracker row
#   OLDER_INCLUDE         completed jobs from the 2024 & 2025 tab to pull into the 2026 import
#   OLDER_DATE_SET        "project|task path" -> date Greg supplied (record-only import)
#   OLDER_CENSUS          project -> [municipality, county, ZIP] from the Census geocoder
#   OLDER_MUNI_CONFIRMED  project -> [municipality, county] Greg confirmed
#   OLDER_SITE            project -> site address Greg supplied
#   OLDER_EQUIPMENT       project -> equipment Greg confirmed
#   OLDER_ANSWERED        "project|flag prefix" -> note for the project log (null drops the flag)
#   OLDER_MUNI_NOTES      project -> municipality note
#   MANUAL_CONTRACTS      contract file name piece -> fields read by hand from the PDF
OVERRIDES_PATH = Path(__file__).resolve().parents[2] / 'import' / 'overrides.json'


def load_overrides(path: Path) -> dict:
	raw = json.loads(path.read_text()) if path.is_file() else {}
	pair_key = lambda d: {tuple(k.split('|', 1)): v for k, v in d.items()}
	tuples = lambda d: {k: tuple(v) for k, v in d.items()}
	return {
		'SKIP_AS_SERVICE': set(raw.get('SKIP_AS_SERVICE', [])),
		'CUSTOMER_ALIASES': [(re.compile(p, re.I), n) for p, n in raw.get('CUSTOMER_ALIASES', [])],
		'CENSUS_MUNI': tuples(raw.get('CENSUS_MUNI', {})),
		'SITE_OVERRIDE': raw.get('SITE_OVERRIDE', {}),
		'CONFIRMED_MUNI': tuples(raw.get('CONFIRMED_MUNI', {})),
		'DATE_FIX': pair_key(raw.get('DATE_FIX', {})),
		'MANUAL_PROJECTS': raw.get('MANUAL_PROJECTS', []),
		'OLDER_INCLUDE': set(raw.get('OLDER_INCLUDE', [])),
		'OLDER_DATE_SET': pair_key(raw.get('OLDER_DATE_SET', {})),
		'OLDER_CENSUS': tuples(raw.get('OLDER_CENSUS', {})),
		'OLDER_MUNI_CONFIRMED': tuples(raw.get('OLDER_MUNI_CONFIRMED', {})),
		'OLDER_SITE': raw.get('OLDER_SITE', {}),
		'OLDER_EQUIPMENT': raw.get('OLDER_EQUIPMENT', {}),
		'OLDER_ANSWERED': pair_key(raw.get('OLDER_ANSWERED', {})),
		'OLDER_MUNI_NOTES': raw.get('OLDER_MUNI_NOTES', {}),
		'MANUAL_CONTRACTS': raw.get('MANUAL_CONTRACTS', {}),
	}


globals().update(load_overrides(OVERRIDES_PATH))


SALES = {'XX': 'person@example.com', 'Staff': 'person@example.com',
		 'Staff': 'person@example.com', 'Staff': 'person@example.com',
		 'GREG': 'person@example.com'}
CUST_TYPE = {'C': ('commercial', False), 'R': ('residential', False), 'A': ('commercial', True),
			 'NP': ('nonprofit', False), 'M': ('government', False)}
INSTALL_TYPE = {'POLE': 'tracker', 'ROOF': 'roof', 'GRND': 'ground', 'GROUND': 'ground'}
UTILITY = {'PPL': 'PPL', 'MET-ED': 'Met-Ed', 'MET - ED': 'Met-Ed', 'METED': 'Met-Ed', 'PENELEC': 'Penelec',
		   'WEST PENN POWER': 'West Penn Power', 'W PENN': 'West Penn Power', 'WPP': 'West Penn Power',
		   'WEST PENN': 'West Penn Power', 'PECO': 'PECO'}
FUNDING = {'REAP': 'REAP Grant', 'ORRSTOWN': 'Orrstown Bank', 'CLIMATIZE': 'Climatize', 'ATMOS': 'Atmos'}
AGENCY = {
	'light-heigel': 'Light-Heigel Associates', 'keystone inspection agency': 'Keystone Inspection Agency',
	'pamca': 'PAMCA', 'technicon': 'Technicon', 'commonwealth': 'Commonwealth', 'kraft municipal': 'Kraft Municipal Group',
	'biu': 'BIU', 'tri-county cog': 'Tri-County COG', 'abi': 'ABI', 'code inspections': 'Code Inspections, Inc',
	'ckcog': 'CKCOG',
}
# AHJ text with no county in the sheet: best guess, always flagged for review
COUNTY_GUESS = {
	'monaghan township': ('York', 'PA'), 'upper leacock township': ('Lancaster', 'PA'),
	'lurgan township': ('Franklin', 'PA'), 'providence township': ('Lancaster', 'PA'),
	'lack township': ('Juniata', 'PA'),
}
AHJ_TYPO = {'palpho township': 'Ralpho Township'}


# Completed jobs on the "2024 & 2025" tab that Greg wants imported. That tab writes some
# numbers short ("42 Sample Farm" = 25042) and has a different column layout from col 8 on.
OLDER_TAB = '2024 & 2025'


def older_number(text: str) -> str | None:
	"""'42 Sample Farm' -> '25042'; '25026 Example Holdings LLC' -> '25026'."""
	m = re.match(r'^(\d{5}|\d{2})\s', text.strip())
	if not m:
		return None
	return m.group(1) if len(m.group(1)) == 5 else '250' + m.group(1)


def remap_older(r) -> tuple[list, str]:
	"""Map a 2024 & 2025 tab row onto the 2026 column layout. Returns (row, email)."""
	r = list(r) + [None] * 40
	new = r[:8] + [None, r[8]] + r[10:26] + [r[26], None, r[27], r[28], r[29], r[31], r[32], r[33], r[34]]
	return new, clean(r[9])


# Spreadsheet column -> task path ("Task" or "Task>Sub-task") in the default template
DATE_COLS = {
	15: 'Deposit received',
	16: 'Design>Planset sent', 17: 'Design>Planset received', 18: 'Design>Engineering stamp received',
	19: 'Zoning permit>Applied', 20: 'Zoning permit>Received',
	21: 'Building permit>Applied', 22: 'Building permit>Received',
	23: 'Interconnection>Applied (app and fee)', 25: 'Interconnection>Approved / permission to install',
	26: 'Installation>Started', 27: 'Installation>Completed',
	28: 'Utility>COC submitted', 29: 'Utility>PTO received',
	30: 'SREC registration', 31: 'SolarInsure registration', 32: 'SolarEdge extended warranty',
}


# --------------------------------------------------------------------------- helpers

def iso(d) -> str | None:
	if isinstance(d, dt.datetime):
		return d.date().isoformat()
	if isinstance(d, dt.date):
		return d.isoformat()
	return None


def clean(v) -> str:
	return '' if v is None else str(v).strip()


def infer_year(month: int, day: int, anchor: dt.date | None) -> dt.date | None:
	"""Pick the year for a month/day with no year: on/after the anchor (contract) date, not far in the future."""
	base = anchor or TODAY
	for y in (base.year, base.year + 1, base.year - 1):
		try:
			d = dt.date(y, month, day)
		except ValueError:
			return None
		if (anchor is None or d >= anchor - dt.timedelta(days=45)) and d <= TODAY + dt.timedelta(days=120):
			return d
	return None


def parse_loose_date(text: str, anchor: dt.date | None) -> dt.date | None:
	m = re.search(r'(\d{1,2})[/-](\d{1,2})(?:[/-](\d{2,4}))?', text)
	if not m:
		return None
	mo, da, yr = int(m.group(1)), int(m.group(2)), m.group(3)
	if yr:
		y = int(yr) + (2000 if len(yr) == 2 else 0)
		try:
			return dt.date(y, mo, da)
		except ValueError:
			return None
	return infer_year(mo, da, anchor)


# --------------------------------------------------------------------------- equipment text

INV_WORDS = re.compile(r'invert|\binv\b|growatt|solaredge|\bse\b|se\d{3,}|home ?hub|hd wave|sma|sbse|chint|fronius|enphase|iq8|peak3|\bkw\b|\d+kw|\bmin \d', re.I)
RACK_WORDS = re.compile(r'\bsats?\b|tracker|sunaction|sun action|mechatron|pole option', re.I)
BATT_WORDS = re.compile(r'batter|energy bank|kwh', re.I)
MICRO_KW = {'IQ8M': 0.325, 'IQ8A': 0.349, 'IQ8PLUS': 0.29, 'IQ8+': 0.29, 'DS3': 0.73}


def split_top(text: str, seps=(',',)) -> list[str]:
	"""Split on separators that are followed by a quantity, ignoring anything inside parentheses."""
	parts, buf, depth, i = [], '', 0, 0
	while i < len(text):
		ch = text[i]
		if ch == '(':
			depth += 1
		elif ch == ')':
			depth = max(0, depth - 1)
		matched = None
		if depth == 0:
			for sep in seps:
				if text.startswith(sep, i):
					rest = text[i + len(sep):].lstrip()
					if re.match(r'\(?\d', rest) and not re.match(r'\d+(\.\d+)?\s*%', rest):
						matched = sep
						break
		if matched:
			parts.append(buf.strip())
			buf = ''
			i += len(matched)
			continue
		buf += ch
		i += 1
	if buf.strip():
		parts.append(buf.strip())
	return [p for p in parts if p]


def lead_qty(part: str) -> tuple[int | None, str]:
	p = re.sub(r'^GM\s+', '', part.strip(), flags=re.I)
	if re.match(r'^1 0 ', p):  # "1 0 11.5 kW" typo for 10
		return 10, p[4:]
	m = re.match(r'^\(?(\d+)\)?\s*(?:-|x)?\s*(.*)$', p)
	return (int(m.group(1)), m.group(2).strip()) if m else (None, p)


def inverter_kw(rest: str) -> float | None:
	for model, kw in MICRO_KW.items():
		if model.lower() in rest.lower().replace(' ', ''):
			return kw
	m = re.search(r'(\d+(?:\.\d+)?)\s*kw(?!h)', rest, re.I)
	if m:
		return float(m.group(1))
	m = re.search(r'\bSE(\d+(?:\.\d+)?)K', rest, re.I)                      # SolarEdge SE100K / SE80KUS
	if m:
		return float(m.group(1))
	m = re.search(r'(?:CORE\d|PEAK\d|TRIPOWER|HIGHPOWER)\s+(\d+(?:\.\d+)?)\b', rest, re.I)   # SMA CORE1 50-US, PEAK3 125-US
	if m:
		return 33.3 if m.group(1) == '33' else float(m.group(1))
	m = re.search(r'(?:SE|MIN\s*|SBSE)?(\d{4,6})(?:H|TL|-)', rest, re.I)
	if m and 1000 <= int(m.group(1)) <= 200000:
		return int(m.group(1)) / 1000
	m = re.search(r'SBSE(\d+(?:\.\d+)?)', rest, re.I)
	if m:
		return float(m.group(1))
	m = re.search(r'(?<![\d.])(\d{4,6})(?![\d.])', rest)          # bare "11400"
	if m and 1000 <= int(m.group(1)) <= 200000:
		return int(m.group(1)) / 1000
	m = re.search(r'(?<![\d.])(\d+(?:\.\d+)?)(?![\d.])', rest)
	if m and 0.2 <= float(m.group(1)) <= 250:
		return float(m.group(1))
	return None


def parse_equipment(text: str, sheet_kw: float | None) -> tuple[dict, list[str]]:
	eq = {'modules': [], 'inverters': [], 'batteries': [], 'racking': None}
	flags: list[str] = []
	t = clean(text)
	if not t or re.search(r'cancel|relocate|r&r', t, re.I):
		if t:
			flags.append(f'Equipment text not parsed: "{t}"')
		return eq, flags

	racking = []
	pending_module: dict | None = None
	for part in split_top(t, (',', ' and ', ' + ', '; ')):
		# parts like "2 - 11.4 kW & 2 - 7.6 kW Growatt" contain several lines
		cls = 'rack' if RACK_WORDS.search(part) else 'batt' if (BATT_WORDS.search(part) and not re.search(r'\bkw\b(?!h)', part, re.I)) \
			else 'inv' if INV_WORDS.search(part) else 'mod'
		if cls == 'rack':
			racking.append(part)
			continue
		sub_parts = split_top(part, (' & ', '& ', ' &')) if '&' in part else [part]
		# trailing words (brand) apply to all sub-parts
		for sp in sub_parts:
			qty, rest = lead_qty(sp)
			if qty is None:
				flags.append(f'Could not read quantity in "{sp}"')
				continue
			if cls == 'inv':
				pairs = re.findall(r'(\d+)\s*x\s*(\d+(?:\.\d+)?)', rest)
				if pairs:
					for q, kw in pairs:
						eq['inverters'].append({'qty': int(q), 'ac_kw': float(kw), 'description': re.sub(r'\(.*?\)', '', rest).strip()})
					continue
				kw = inverter_kw(rest)
				if kw is None:
					flags.append(f'Inverter size unknown in "{sp}" (line left out)')
					continue
				desc = re.sub(r'\b\d+(?:\.\d+)?\s*kw\b', '', rest, flags=re.I).strip(' -,')
				eq['inverters'].append({'qty': qty, 'ac_kw': kw, 'description': desc or None})
			elif cls == 'batt':
				kwh = re.search(r'(\d+(?:\.\d+)?)\s*kwh', rest, re.I)
				eq['batteries'].append({'qty': qty, 'kwh': float(kwh.group(1)) if kwh else None, 'kw': None,
										'description': rest})
			else:
				w = re.search(r'(?<!\d)([3-7]\d{2})(?:\s*(?:w|watt)\b)?', rest, re.I)
				if w:
					desc = re.sub(r'\b' + w.group(1) + r'\s*(?:w|watt)\b', w.group(1) + 'W', rest, count=1, flags=re.I)
					pending_module = {'qty': qty, 'watts': float(w.group(1)), 'description': desc.strip(' -')}
					eq['modules'].append(pending_module)
				elif pending_module is not None:
					# "369 SEG 590, 332 Grade B, 37 US": split the total into sub-lines
					eq['modules'].append({'qty': qty, 'watts': pending_module['watts'],
										  'description': (pending_module['description'] + ' ' + rest).strip(), '_sub_of': id(pending_module)})
				else:
					eq['modules'].append({'qty': qty, 'watts': None, 'description': rest})
	# collapse "total + sub-lines"
	totals = {id(m): m for m in eq['modules'] if 'watts' in m}
	subs: dict[int, list] = {}
	for m in eq['modules']:
		if '_sub_of' in m:
			subs.setdefault(m['_sub_of'], []).append(m)
	for tid, lines in subs.items():
		total = totals.get(tid)
		if total and sum(l['qty'] for l in lines) == total['qty']:
			eq['modules'].remove(total)
		elif total:
			flags.append(f"Module sub-lines don't add up to {total['qty']} in \"{t}\"")
	for m in eq['modules']:
		m.pop('_sub_of', None)
		base = re.sub(r'\b(panels?)\b', '', m['description'] or '', flags=re.I)
		m['description'] = re.sub(r'\s+', ' ', base).strip() or None
	# watts unknown: derive from sheet kW if it's the only module line
	for m in eq['modules']:
		if m['watts'] is None:
			if sheet_kw and len(eq['modules']) == 1 and m['qty']:
				m['watts'] = round(sheet_kw * 1000 / m['qty'])
				flags.append(f"Module wattage not in sheet; used {m['watts']}W from kW column")
			else:
				flags.append(f"Module wattage unknown for \"{m['description']}\" (line left out)")
	eq['modules'] = [m for m in eq['modules'] if m['watts']]
	if racking:
		eq['racking'] = '; '.join(racking)
	dc = sum(m['qty'] * m['watts'] for m in eq['modules']) / 1000
	if sheet_kw and dc and abs(dc - sheet_kw) / sheet_kw > 0.01:
		flags.append(f'Parsed DC {dc:.2f} kW vs sheet {sheet_kw:g} kW')
	return eq, flags


# --------------------------------------------------------------------------- municipality

def parse_ahj(text: str) -> tuple[dict | None, list[str], str | None]:
	"""Return ({name, county, state}, flags, agency-from-parentheses)."""
	t = clean(text)
	flags: list[str] = []
	if not t:
		return None, flags, None
	agency = None
	m = re.search(r'\((CKCOG|[A-Z]{2,}COG)\)', t)
	if m:
		agency = m.group(1)
		t = t.replace(m.group(0), '').strip(' ,')
	county = None
	# "Lebanon County (AHJ), South Annville Township"
	m = re.match(r'^(\w[\w ]*?) County \(AHJ\),\s*(.+)$', t)
	if m:
		county, t = m.group(1), m.group(2)
		flags.append(f'Sheet says {county} County is the AHJ for {t}; recorded the township, add the county office as the agency if needed')
	m = re.match(r'^(.+?)\s*\((\w[\w ]*?)(?: County)?\)$', t)          # "Penn Twp (Lancaster)"
	if m and not county:
		t, county = m.group(1), m.group(2)
	m = re.match(r'^(.+?)[,\s]+([A-Z][a-z]+) County$', t)               # "X Twp, Y County" / "X Township Y County"
	if m and not county:
		t, county = m.group(1), m.group(2)
	m = re.match(r'^(.+?),\s*(\w[\w ]*?)$', t)                          # "Manheim Township, Lancaster"
	if m and not county:
		t, county = m.group(1), m.group(2)
	name = re.sub(r'\bTwp\.?\b', 'Township', t.strip(' ,'), flags=re.I)
	name = re.sub(r'\bTP\b', 'Township', name)
	if not re.search(r'township|borough|city|town\b', name, re.I):
		name += ' Township'
		flags.append(f'Assumed "{name}" is a township')
	key = name.lower()
	if key in AHJ_TYPO:
		flags.append(f'Sheet spelling "{name}" corrected to "{AHJ_TYPO[key]}"')
		name = AHJ_TYPO[key]
		key = name.lower()
	state = 'PA'
	if not county:
		if key in COUNTY_GUESS:
			county, state = COUNTY_GUESS[key]
			flags.append(f'County not in sheet; guessed {county} County for {name}')
		else:
			flags.append(f'County unknown for "{name}"; municipality left blank')
			return None, flags, agency
	return {'name': name, 'county': county.strip(), 'state': state}, flags, agency


def parse_agency(text: str) -> tuple[str | None, list[str]]:
	t = clean(text)
	if not t:
		return None, []
	low = t.lower()
	flags = []
	if '&' in t or low.startswith('(2)'):
		flags.append(f'Multiple inspection agencies listed: "{t}"; used the first')
	if '?' in t:
		flags.append(f'Inspection agency marked uncertain: "{t}"')
	for key, name in AGENCY.items():
		if key in low:
			return name, flags
	return re.sub(r'[?]', '', t).strip(), flags


# --------------------------------------------------------------------------- contracts

def pdf_text(path: Path) -> tuple[str, bool]:
	"""Text of a PDF; scanned PDFs (no text layer) are OCR'd with tesseract. Returns (text, used_ocr)."""
	text = subprocess.run(['pdftotext', '-layout', str(path), '-'], capture_output=True, text=True, check=True).stdout
	if len(text.split()) >= 200:
		return text, False
	import tempfile
	with tempfile.TemporaryDirectory() as tmp:
		subprocess.run(['pdftoppm', '-r', '200', '-png', str(path), f'{tmp}/p'], check=True)
		pages = sorted(Path(tmp).glob('p-*.png'))
		out = [subprocess.run(['tesseract', str(pg), 'stdout', '--psm', '6'], capture_output=True, text=True).stdout for pg in pages]
	text = '\n'.join(out)
	text = re.sub(r'(\d)\.x ', r'\1 x ', text)           # OCR: "3.x MIN" -> "3 x MIN"
	return text, True


STREET_SUFFIX = r'(?:Ln|Lane|Rd|Road|Dr|Drive|Ave|Avenue|St|Street|Blvd|Boulevard|Way|Pike|Ct|Court|Cir|Circle|Hwy|Pl|Place|Ter|Trl|Run)\.?'


def split_street_city(s: str) -> tuple[str, str | None]:
	"""'100 N Main St, Springville' or '12 Orchard Ln Honey Brook,' -> (street, city)."""
	s = s.strip(' ,')
	if ',' in s:
		street, city = (x.strip() for x in s.rsplit(',', 1))
		if city and not re.match(r'^(PA|MD|DE)\b', city):
			return street, city
		s = street
	m = re.match(r'^(\d+\s.*?\b' + STREET_SUFFIX + r')\s+([A-Z][A-Za-z .]+)$', s)
	return (m.group(1), m.group(2).strip()) if m else (s, None)


def unkern(s: str) -> str:
	"""pdftotext sometimes splits a capital from the rest of the word: 'N ew Holland', 'W illiam'."""
	return re.sub(r'\b([A-Z]) ([a-z]{2,})\b', r'\1\2', s)


def parse_contract(path: Path) -> dict:
	raw, ocr = pdf_text(path)
	text = raw.replace('\u202d', '').replace('\u202c', '')
	for lig, rep in {'\ufb00': 'ff', '\ufb01': 'fi', '\ufb02': 'fl', '\ufb03': 'ffi', '\ufb04': 'ffl'}.items():
		text = text.replace(lig, rep)
	lines = [l for l in text.splitlines() if l.strip()]
	c: dict = {'file': path.name, 'flags': ['Scanned contract read by OCR: double-check its numbers'] if ocr else []}

	m = re.search(r'Prepared by:?\s*([A-Z][a-z]+ [A-Z][a-z]+)', text)
	c['salesperson'] = m.group(1) if m else None
	m = re.search(r'For\s*:?\s+([^\n]+?)(?:\s{2,}|\n)', text)
	c['contact_name'] = re.sub(r'\s*Quote\b.*$', '', m.group(1)).strip(' :') if m else None
	if c['contact_name'] in (None, ''):
		m = re.search(r'For\s+([A-Z][^\n]+?)\s*\n', text)
		c['contact_name'] = m.group(1).strip() if m else None
	m = re.search(r'Quote #:\s*(\d+)', text)
	c['quote_number'] = m.group(1) if m else None

	# Address block: middle column of the header lines
	mid = []
	for l in lines[:8]:
		cols = [x.strip() for x in re.split(r'\s{2,}', l.strip()) if x.strip()]
		if len(cols) >= 2:
			mid.append(cols[1] if not cols[0].startswith(('For', 'Prepared')) else (cols[1] if len(cols) > 1 else ''))
		elif len(cols) == 1 and l.startswith(' ' * 20):
			mid.append(cols[0])
	c['email'] = next((x for x in mid if '@' in x and 'trifecta' not in x.lower()), None)
	phone = next((x for x in mid if re.fullmatch(r'[\d\s()\-.+]{10,}', x)), None)
	c['phone'] = re.sub(r'^\(1\)\s*', '', phone) if phone else None
	street = city = state = zipc = None
	for x in mid:
		z = re.search(r'^(.*?),?\s*\b(PA|MD|DE)\s+(\d{5})', x)
		if z:
			state, zipc = z.group(2), z.group(3)
			before = z.group(1).strip(' ,')
			if re.match(r'^\d+\s', before):        # "34 Mill Rd Lititz, PA 17543": street and city on one line
				street, city = split_street_city(before)
			else:
				city = before or None
	if not street:
		street = next((x for x in mid if re.match(r'^\d+\s+\w', x) and not re.search(r'Quote|Valid', x)), None)
		if street and not city:
			street, city = split_street_city(street)  # "12 Orchard Ln Honey Brook," / "100 N Main St, Springville"
	if street and not city:
		c['flags'].append('Contract address has no separate city line; check the site address')
	city = unkern(city) if city else city
	if c.get('contact_name'):
		c['contact_name'] = unkern(c['contact_name'])
	# pdftotext drops some ligature glyphs ("Mi in" for Mifflin); repair the common PA ones
	city_fix = {'Mi in': 'Mifflin', 'Mi intown': 'Mifflintown', 'Mi inburg': 'Mifflinburg', 'Gri n': 'Griffin'}
	if city in city_fix:
		city = city_fix[city]
	c['address'] = {'street': street, 'city': city, 'state': state, 'zip': zipc}

	m = re.search(r'Total System Price\s+\$([\d,]+(?:\.\d\d)?)', text)
	c['price'] = float(m.group(1).replace(',', '')) if m else None
	m = re.search(r'([\d,]+) kWh per year', text)
	c['annual_kwh'] = int(m.group(1).replace(',', '')) if m else None
	m = re.search(r'savings calculated based on\s+(.+?)\s+electric\s*rate', text, re.S)
	c['rate_note'] = re.sub(r'\s+', ' ', m.group(1)) if m else None
	c['solarinsure'] = bool(re.search(r'Solar\s*Insure', text, re.I) and re.search(r'Includes Solar Insure|Included in the contract price is the Solar Insure', re.sub(r'\s+', ' ', text), re.I))
	m = re.search(r'Payment Option:\s*([^\n]+)', text)
	c['payment_option'] = m.group(1).strip() if m else None

	# Quotation lines
	q = text[text.find('Payment Option'):] if 'Payment Option' in text else text
	c['modules'] = []
	for m in re.finditer(r'^\s*(\d+) x (.+?)\s(\d{3}(?:\.\d)?) Watt Panels?(?: \((.+?)\))?\s*$', q, re.M):
		c['modules'].append({'qty': int(m.group(1)), 'watts': float(m.group(3)),
							 'description': (m.group(2) + (' ' + m.group(4) if m.group(4) else '')).strip()})
	c['inverters'] = []
	c['batteries'] = []
	c['racking'] = []
	c['other_items'] = []
	batt_total = re.search(r'([\d.]+)\s*kWh Total Battery Storage', text)
	section = q[:q.find('Total System Price')] if 'Total System Price' in q else q
	for line in section.splitlines():
		line = line.strip()
		if not re.match(r'^\d+ x ', line) or re.search(r'Watt Panels?', line):
			continue
		brand_m = re.search(r'\(([^()]*)\)\s*$', line)
		brand = brand_m.group(1).split()[0] if brand_m and not re.search(r'Grade|US Content', brand_m.group(1)) else ''
		body = line[:brand_m.start()].strip() if brand else line
		for item in re.split(r',\s*(?=\d+ x )', body):
			m = re.match(r'^(\d+) x (.+)$', item.strip())
			if not m:
				continue
			qty, desc = int(m.group(1)), m.group(2).strip()
			if re.search(r'optimi[sz]er|\bRSD\b|^[PSU]\d{3}[A-Z]?\b|^C\d{3}U?\b|APSM', desc, re.I):
				continue
			if RACK_WORDS.search(desc):
				c['racking'].append(f'{qty} x {desc}')
				continue
			if re.search(r'transformer|kva|monitoring|cellular', desc, re.I):
				c['other_items'].append(f'{qty} x {desc}')
				continue
			desc = desc.replace('[240VI', '[240V]')
			if re.search(r'\bBAT-|batter', desc, re.I):
				kwh = float(batt_total.group(1)) / qty if batt_total else None
				c['batteries'].append({'qty': qty, 'kwh': kwh, 'kw': None, 'description': (brand + ' ' + desc).strip()})
				continue
			kw = inverter_kw(desc)
			if kw is None:
				c['flags'].append(f'Contract line not understood: "{qty} x {desc}"')
				continue
			c['inverters'].append({'qty': qty, 'ac_kw': kw, 'description': (brand + ' ' + desc).strip()})

	# Payment milestones (labels only, no dollar amounts)
	c['payments'] = []
	pm = text.find('Payment Milestones')
	if pm >= 0:
		# One row per amount line. A long label can wrap: then the amount line has no text and
		# the label is the text just above it plus any continuation right below it.
		body = []
		for l in text[pm:].splitlines()[1:]:
			if re.match(r'^\s*Total\b', l):
				break
			body.append(l)
		amount_at_end = re.compile(r'^\s*(.*?)\s*\|?\s*\$[\d,]+(?:\.\d\d)?\s*$')
		is_amount = lambda ln: bool(amount_at_end.match(ln))
		used = set()
		for i, l in enumerate(body):
			mm = amount_at_end.match(l)
			if not mm:
				continue
			label = mm.group(1).strip(' |')
			if not label:
				above, j = [], i - 1
				while j >= 0 and j not in used and body[j].strip() and not is_amount(body[j]):
					above.insert(0, body[j].strip()); used.add(j); j -= 1
				below, j = [], i + 1
				while j < len(body) and body[j].strip() and not is_amount(body[j]):
					below.append(body[j].strip()); used.add(j); j += 1
				label = ' '.join(above + below)
			label = re.sub(r'^[I|]\s+', '', re.sub(r'\s+', ' ', label)).strip()
			if label:
				c['payments'].append(label)

	# Signed date: e-sign certificate, else a trailing handwritten-style date
	m = re.search(r'eSigned by:\s*\n?\s*(\d{2}/\d{2}/\d{4})', text) or re.search(r'Document completed by all parties on (\d{2}/\d{2}/\d{4})', text)
	if m:
		c['signed'] = dt.datetime.strptime(m.group(1), '%m/%d/%Y').date().isoformat()
	else:
		m = re.findall(r'^\s*(\d{1,2}-\d{1,2}-\d{2})\s*$', text, re.M)
		c['signed'] = dt.datetime.strptime(m[-1], '%m-%d-%y').date().isoformat() if m else None
		if c['signed']:
			c['flags'].append('Signed date read from a hand-entered date on the contract')
	return c


def name_tokens(s: str) -> set[str]:
	stop = {'inc', 'llc', 'the', 'contract', 'signed', 'revised', 'farm', 'farms', 'and', 'of', 'pdf', 'proposal',
			'merged', 'final', 'panel', 'panels', 'seg', 'grade', 'pole', 'mount', 'company', 'co', 'tp', 'nd', 'rd',
			'june', 'growatt', 'inverters', 'k', 'a'}
	return {w for w in re.findall(r'[a-z]+', s.lower()) if w not in stop and len(w) > 1}


# --------------------------------------------------------------------------- main

def main() -> int:
	ap = argparse.ArgumentParser()
	ap.add_argument('--xlsx', required=True)
	ap.add_argument('--sheet', default='2026')
	ap.add_argument('--drive')
	ap.add_argument('--contracts', nargs='*', default=[])
	ap.add_argument('--out', required=True)
	ap.add_argument('--report', required=True)
	a = ap.parse_args()

	wb = openpyxl.load_workbook(a.xlsx, data_only=True)
	ws = wb[a.sheet]

	drive: dict[str, dict] = {}
	drive_unmatched = []
	if a.drive:
		with open(a.drive, newline='', encoding='utf-8-sig') as fh:
			for row in csv.DictReader(fh, delimiter='\t'):
				m = re.match(r'^(\d{5})\b', row['Folder Name'].strip())
				short = older_number(row['Folder Name'])
				if not m and short in OLDER_INCLUDE:
					m = re.match(r'^(\d+)', short)
				if m:
					drive[m.group(1)] = {'url': row['URL'].strip(), 'folder': row['Folder Name'].strip(), 'location': row['Location'].strip()}

	contracts = []
	for p in a.contracts:
		path = Path(p)
		for f in sorted(path.glob('*.pdf') if path.is_dir() else [path]):
			try:
				contracts.append(parse_contract(f))
			except Exception as e:  # keep going; report it
				contracts.append({'file': f.name, 'error': str(e), 'flags': [f'Could not read contract: {e}']})

	projects, skipped, report_rows = [], [], []
	solarinsure_numeric: list[str] = []
	seen = set()
	rows = []
	for r in ws.iter_rows(min_row=4, values_only=True):
		if clean(r[0]).startswith('2024'):
			break
		rows.append((r, False, ''))
	if OLDER_TAB in wb.sheetnames:
		for r in wb[OLDER_TAB].iter_rows(min_row=4, values_only=True):
			num = older_number(clean(r[0]))
			if num in OLDER_INCLUDE:
				nr, email = remap_older(r)
				nr[0] = re.sub(r'^\d+', num, clean(r[0]))
				rows.append((nr, True, email))
	for r, is_older, sheet_email in rows:
		name_cell = clean(r[0])
		if not name_cell:
			continue
		m = re.match(r'^(\d{5})\s+(.+)$', name_cell)
		if not m:
			skipped.append((name_cell, 'No 5-digit project # at the start of the name'))
			continue
		number, pname = m.group(1), re.sub(r'\\', '', m.group(2)).strip()
		if number in SKIP_AS_SERVICE:
			skipped.append((name_cell, 'Re-roof / repair work: belongs in Service, not imported as a project'))
			continue
		seen.add(number)
		flags: list[str] = []
		signed = iso(r[1])
		if not signed and clean(r[1]):
			d = parse_loose_date(clean(r[1]), None)
			signed = d.isoformat() if d else None
			if not signed:
				flags.append(f'Contract signed cell is "{clean(r[1])}", not a date')
		anchor = dt.date.fromisoformat(signed) if signed else None

		sales = SALES.get(clean(r[2]).upper())
		if clean(r[2]) and not sales:
			flags.append(f'Unknown salesperson "{clean(r[2])}"')
		ctype, ag = CUST_TYPE.get(clean(r[4]).upper(), (None, False))
		sheet_kw = float(r[3]) if isinstance(r[3], (int, float)) else None

		status_raw = clean(r[14])
		hold_state = hold_reason = None
		status_note = status_raw or None
		if re.match(r'^(paused|on hold)\b', status_raw, re.I):
			hold_state = 'on_hold'
			hold_reason = re.sub(r'^(paused|on hold)\s*(-|for)?\s*', '', status_raw, flags=re.I).strip() or None
		elif re.search(r'project cancelled', status_raw, re.I):
			hold_state, hold_reason = 'cancelled', status_raw
		if re.search(r'project cancelled', clean(r[9]), re.I) and hold_state != 'cancelled':
			hold_state, hold_reason = 'cancelled', 'Project cancelled (equipment column)'
			if status_raw:
				status_note = status_raw

		eq, eq_flags = parse_equipment(clean(r[9]), sheet_kw)
		flags += eq_flags
		muni, muni_flags, muni_agency = parse_ahj(clean(r[11]))
		flags += muni_flags
		agency, ag_flags = parse_agency(clean(r[12]))
		flags += ag_flags
		agency = agency or muni_agency

		util = UTILITY.get(clean(r[10]).upper().replace('  ', ' '))
		if clean(r[10]) and not util:
			flags.append(f'Unknown utility "{clean(r[10])}"')
		funding = [FUNDING[k] for k in FUNDING if k in clean(r[7]).upper()]
		tax = {'YES': 1, 'NO': 0}.get(clean(r[6]).upper())

		# ---- tasks
		tasks: dict[str, dict] = {}

		def put(path, **kw):
			tasks.setdefault(path, {}).update({k: v for k, v in kw.items() if v is not None})

		for col, path in DATE_COLS.items():
			v = r[col]
			if v is None or clean(v) == '':
				continue
			d = iso(v)
			s = clean(v)
			if d:
				if d > (TODAY + dt.timedelta(days=0 if is_older else 60)).isoformat():
					flags.append(f'{path}: {d} is in the future (left as-is, fix if the year is wrong)')
				put(path, done=d, needed=1)
			elif s.upper() in ('N/A', '--', 'NA'):
				put(path, needed=0)
			elif path == 'Design>Engineering stamp received' and s.upper() == 'NO':
				put(path, needed=0, note='No stamp required (spreadsheet)')
			elif path == 'SolarInsure registration' and re.fullmatch(r'\d+(\.0)?', s):
				# "10" = no SolarInsure; Trifecta's own 10-year warranty applies
				put(path, needed=0, note=f'No SolarInsure (spreadsheet "{s.split(".")[0]}": Trifecta {s.split(".")[0]}-year warranty)')
				solarinsure_numeric.append(number)
			elif s.upper().startswith('YES') or s.lower() in ('done', 'missing items'):
				dd = parse_loose_date(s, anchor)
				if dd:
					put(path, needed=1, done=dd.isoformat(), note=f'Spreadsheet: "{s}"')
				elif s.lower() == 'missing items':
					put(path, needed=1, note='Spreadsheet: missing items')
				else:
					proxy = signed or None
					put(path, needed=1, done=proxy, note='Spreadsheet said Yes; actual date unknown' + ('' if proxy else ' (no contract date either)'))
			else:
				dd = parse_loose_date(s, anchor)
				if dd and re.search(r'\d+/\d+', s):
					put(path, needed=1, done=dd.isoformat(), note=f'Spreadsheet: "{s}" (year assumed)')
					if re.search(r'A&J', s):
						flags.append(f'{path}: "{s}" suggests A&J did the work; set the Installer if so')
				else:
					put(path, note=f'Spreadsheet: "{s}"')

		# Zoning: both N/A -> whole task not needed
		if tasks.get('Zoning permit>Applied', {}).get('needed') == 0 and tasks.get('Zoning permit>Received', {}).get('needed') == 0:
			put('Zoning permit', needed=0)
		# Utility change required
		uc = clean(r[24])
		if uc:
			if uc.upper() in ('NO', 'N/A', '--'):
				put('Utility upgrade', needed=0)
			elif uc.upper() == 'YES' or 'upgrade' in uc.lower():
				put('Utility upgrade', needed=1, note=None if uc.upper() == 'YES' else f'Spreadsheet: "{uc}"')
			else:
				put('Utility upgrade', note=f'Spreadsheet: "{uc}"')
		ic = clean(r[25])
		if ic.lower().startswith('cond'):
			tasks.pop('Interconnection>Approved / permission to install', None)
			put('Interconnection>Conditional approval', needed=1, note='Spreadsheet: "cond." (conditional approval, date unknown)')
			flags.append('Interconnection shows conditional approval with no date; enter the date to make it count toward Clear to install')
		if clean(r[8]).upper() == 'YES':
			put('Utility rebate', needed=1)
		vnm = clean(r[13]).upper()
		if vnm in ('YES', 'NO'):
			put('Virtual net metering', needed=1 if vnm == 'YES' else 0)
		if 'Deposit received' in tasks and tasks['Deposit received'].get('note', '').startswith('Spreadsheet said Yes'):
			tasks['Deposit received']['note'] = 'Deposit marked received in spreadsheet; date unknown, contract date used'

		dr = drive.get(number)
		customer = next((alias for rx, alias in CUSTOMER_ALIASES if rx.search(pname)), None) or re.sub(r'\s*\(.*?\)\s*', ' ', pname).strip()

		proj = {
			'number': number, 'name': pname, 'customer': customer, 'signed': signed,
			'salesperson_email': sales, 'customer_type': ctype, 'is_agricultural': ag,
			'install_type': INSTALL_TYPE.get(clean(r[5]).upper()), 'sheet_kw': sheet_kw,
			'tax_exempt': tax, 'funding': funding, 'utility': util, 'municipality': muni,
			'inspection_agency': agency, 'status_note': status_note, 'hold_state': hold_state, 'hold_reason': hold_reason,
			'drive_url': dr['url'] if dr else None,
			'archived': bool(dr and dr['location'].lower().startswith('archive')),
			'has_batteries': bool(eq['batteries']), 'has_pv': bool(eq['modules']) or not eq['batteries'],
			'equipment': eq, 'equipment_text': clean(r[9]) or None, 'tasks': tasks,
			'contact': None, 'site': None, 'price': None, 'annual_kwh': None, 'quote_number': None, 'payments': None,
			'source_notes': [], 'flags': flags,
		}
		if is_older:
			proj['archived'] = True
			proj['source_notes'].append(f'From the {OLDER_TAB} tab; imported as Completed')
			proj['_sheet_email'] = sheet_email or None
			if r[26] or r[29]:
				put('Installation>Completed', needed=1, note=f'Install complete date not tracked on the {OLDER_TAB} tab')
		if not dr:
			flags.append('No Google Drive folder found for this project #')
		if proj['archived'] and dr and not is_older:
			proj['source_notes'].append(f"Drive folder is in {dr['location']}; imported as Completed")
		projects.append(proj)

	# ---- merge contracts (greedy best match, one contract per project)
	cand = []
	for ci, c in enumerate(contracts):
		if 'error' in c:
			report_rows.append(('?', c['file'], c['flags']))
			continue
		ctoks = name_tokens(re.sub(r'\(.*?\)', '', c['file'].rsplit('.', 1)[0])) | name_tokens(c.get('contact_name') or '')
		for p in projects:
			ptoks = name_tokens(p['name'] + ' ' + p['customer'])
			sc = len(ctoks & ptoks)
			if sc:
				cand.append((sc, -len(ptoks - ctoks), ci, p['number']))
	cand.sort(reverse=True)
	assigned_c, assigned_p, matches = set(), set(), []
	for sc, _, ci, num in cand:
		if ci in assigned_c or num in assigned_p:
			continue
		assigned_c.add(ci); assigned_p.add(num); matches.append((ci, num))
	by_num = {p['number']: p for p in projects}
	for ci, c in enumerate(contracts):
		if 'error' not in c and ci not in assigned_c:
			report_rows.append(('?', c['file'], ['Matched no 2026-tab project (older project or name mismatch)']))
	for ci, num in matches:
		c = contracts[ci]
		best = by_num[num]
		p = best
		p['source_notes'].append(f"Contract: {c['file']}")
		p['contract_file'] = c['file']
		p['flags'] += [f'Contract: {f}' for f in c['flags']]
		p['contact'] = {'name': c.get('contact_name'), 'email': c.get('email'), 'phone': c.get('phone')}
		p['site'] = c['address']
		p['price'] = c.get('price')
		p['annual_kwh'] = c.get('annual_kwh')
		p['quote_number'] = c.get('quote_number')
		if c.get('rate_note'):
			p['source_notes'].append(f"Contract savings basis: {c['rate_note']}")
		for item in c.get('other_items', []):
			p['source_notes'].append(f'Contract also includes: {item}')
		if c.get('payment_option'):
			p['source_notes'].append(f"Payment option: {c['payment_option']}")
		if c.get('payments'):
			p['payments'] = c['payments']
		si = p['tasks'].setdefault('SolarInsure registration', {})
		if c.get('solarinsure'):
			if si.get('needed') == 0:
				p['flags'].append('SolarInsure: spreadsheet says no (10), contract includes it (used contract)')
				si.pop('note', None)
			si['needed'] = 1
		elif 'needed' not in si:
			si['needed'] = 0
			si['note'] = 'Not in contract'
		if c.get('signed') and p['signed'] and c['signed'] != p['signed']:
			gap = (dt.date.fromisoformat(c['signed']) - dt.date.fromisoformat(p['signed'])).days
			if 0 < gap <= 30:
				p['flags'].append(f"Signed date: spreadsheet {p['signed']}, contract e-signed {c['signed']} (used contract)")
				p['signed'] = c['signed']
			else:
				p['flags'].append(f"Signed date: spreadsheet {p['signed']}, contract {c['signed']} (kept spreadsheet; contract is likely a revision)")
		elif c.get('signed') and not p['signed']:
			p['signed'] = c['signed']
		# Equipment from the contract is authoritative when present
		old = p['equipment']
		if c['modules']:
			sheet_desc = '; '.join(f"{m['qty']}x{m['watts']:g}W" for m in old['modules'])
			new_desc = '; '.join(f"{m['qty']}x{m['watts']:g}W" for m in c['modules'])
			# keep the sheet's Grade B / US split when the totals and wattage agree
			same_total = sum(m['qty'] for m in old['modules']) == sum(m['qty'] for m in c['modules']) \
				and {m['watts'] for m in old['modules']} == {m['watts'] for m in c['modules']}
			if not same_total:
				p['flags'].append(f'Modules: spreadsheet {sheet_desc or "none"} vs contract {new_desc} (used contract)')
				old['modules'] = c['modules']
		if c['inverters']:
			sd = '; '.join(sorted(f"{i['qty']}x{i['ac_kw']:g}kW" for i in old['inverters']))
			nd = '; '.join(sorted(f"{i['qty']}x{i['ac_kw']:g}kW" for i in c['inverters']))
			if sd != nd:
				p['flags'].append(f'Inverters: spreadsheet {sd or "none"} vs contract {nd} (used contract)')
			old['inverters'] = c['inverters']
		if c['racking']:
			old['racking'] = '; '.join(c['racking'])
		sp = c.get('salesperson')
		if sp and p['salesperson_email'] and sp.split()[0].lower() not in p['salesperson_email']:
			p['flags'].append(f'Salesperson: spreadsheet {p["salesperson_email"]}, contract {sp}')

	for num, d in drive.items():
		if num not in seen and num[:2] in ('25', '26') and num not in SKIP_AS_SERVICE:
			drive_unmatched.append((num, d['folder'], d['location']))

	# Hand-added projects (Drive folder only)
	for mp in MANUAL_PROJECTS:
		if mp['number'] in seen:
			continue
		d = drive.get(mp['number'], {})
		projects.append({
			'number': mp['number'], 'name': mp['name'], 'customer': mp['customer'], 'signed': None,
			'salesperson_email': None, 'customer_type': None, 'is_agricultural': False, 'install_type': None,
			'sheet_kw': None, 'tax_exempt': None, 'funding': [], 'utility': None, 'municipality': None,
			'inspection_agency': None, 'status_note': None, 'hold_state': None, 'hold_reason': None,
			'drive_url': d.get('url'), 'archived': False, 'has_batteries': False, 'has_pv': True,
			'equipment': {'modules': [], 'inverters': [], 'batteries': [], 'racking': None},
			'equipment_text': None, 'tasks': {}, 'contact': None, 'site': None, 'price': None,
			'annual_kwh': None, 'quote_number': None, 'payments': None,
			'source_notes': [mp['note']], 'flags': ['Added by hand: fill in address, municipality, signed date and equipment'],
		})
		seen.add(mp['number'])
		drive_unmatched[:] = [x for x in drive_unmatched if x[0] != mp['number']]

	for p in projects:
		if p['number'] in SITE_OVERRIDE and not (p.get('site') or {}).get('street'):
			p['site'] = SITE_OVERRIDE[p['number']]
			p['source_notes'].append('Site address from Greg')
		if p['number'] == '26019':
			p['flags'].append('Letter of intent, not a final contract (to become a contract by 2/28/26)')
		for (num, path), d in DATE_FIX.items():
			if p['number'] == num and path in p['tasks']:
				p['tasks'][path]['done'] = d
				p['flags'] = [f for f in p['flags'] if not f.startswith(path + ':')]
				p['source_notes'].append(f'{path}: date corrected to {d} (Greg)')
		email = p.pop('_sheet_email', None)
		if email and not (p.get('contact') or {}).get('email'):
			p['contact'] = {**(p.get('contact') or {'name': p['customer'], 'phone': None}), 'email': email}

	# Census address check for the municipality (or Greg's confirmation where Census had no match)
	for p in projects:
		if p['number'] in CONFIRMED_MUNI:
			n, co = CONFIRMED_MUNI[p['number']]
			p['flags'] = [f for f in p['flags'] if not f.startswith(('County not in sheet', 'County unknown'))]
			p['source_notes'].append('Municipality confirmed by Greg (project notes)')
			p['municipality'] = {'name': n, 'county': co, 'state': 'PA'}
			continue
		c = CENSUS_MUNI.get(p['number'])
		if not c:
			continue
		m = p.get('municipality') or {}
		if (m.get('name', '').lower(), m.get('county', '').lower()) == (c[0].lower(), c[1].lower()):
			p['flags'] = [f for f in p['flags'] if not f.startswith('County not in sheet')]
			p['source_notes'].append('Municipality confirmed by Census address lookup')
		else:
			was = f"{m['name']}, {m['county']} Co." if m else 'blank'
			p['flags'] = [f for f in p['flags'] if not f.startswith(('County not in sheet', 'County unknown'))]
			p['flags'].append(f'Municipality set to {c[0]}, {c[1]} Co. from Census address lookup (sheet: {was})')
			p['municipality'] = {'name': c[0], 'county': c[1], 'state': 'PA'}

	out = {'generated': dt.datetime.now().isoformat(timespec='seconds'), 'source': Path(a.xlsx).name, 'projects': projects}
	Path(a.out).parent.mkdir(parents=True, exist_ok=True)
	Path(a.out).write_text(json.dumps(out, indent=1, default=str))

	# ---- report
	L = [f"# Import review: {len(projects)} projects", '',
		 f"Source: {Path(a.xlsx).name} ({a.sheet} tab), {len(drive)} Drive folders, {len(contracts)} contract(s). Generated {out['generated']}.", '',
		 'Everything below was imported, but is a guess or a conflict worth a look. Projects with no notes are omitted.', '']
	if skipped:
		L += ['## Not imported', ''] + [f'- **{n}**: {why}' for n, why in skipped] + ['']
	if drive_unmatched:
		L += ['## Drive folders with no tracker row', ''] + [f'- {n} {f} ({loc})' for n, f, loc in drive_unmatched] + ['']
	if solarinsure_numeric:
		L += [f'SolarInsure column "10" (no SolarInsure, Trifecta 10-year warranty) on {len(solarinsure_numeric)} projects: marked SolarInsure not needed, unless the contract includes it.', '']
	missing = [f"{p['number']} {p['name']}" for p in projects if not p.get('contract_file') and p['hold_state'] != 'cancelled']
	if missing:
		L += ['## Projects with no contract', ''] + [f'- {m}' for m in missing] + ['']
	if report_rows:
		L += ['## Contracts', ''] + [f'- {f}: {"; ".join(fl)}' for _, f, fl in report_rows] + ['']
	L += ['## Per project', '']
	for p in projects:
		notes = p['flags'] + [n for n in p['source_notes'] if n.startswith('Contract:')]
		if not notes:
			continue
		L.append(f"### {p['number']} {p['name']}")
		L += [f'- {f}' for f in notes]
		L.append('')
	Path(a.report).write_text('\n'.join(L))
	print(f"{len(projects)} projects, {len(skipped)} skipped, {sum(1 for p in projects if p['flags'])} with review notes -> {a.out}, {a.report}")
	return 0


# --------------------------------------------------------------------------- record-only import of the older tab
#
# Older jobs come in so reports have history. Only the dates on the sheet become task rows
# (the importer drops every other template step and all payments), PTO present means
# Completed, cancelled rows come in as Cancelled, and project numbers that already exist
# are skipped. Contracts fill what the sheet does not have (price, address, contact, kWh)
# and supply equipment lines when their DC total matches the sheet's kW.

# Two-digit numbers kept as-is (early jobs had no year code); other two-digit 2025 rows are 250xx.
OLDER_KEEP_SHORT = {f'{n:02d}' for n in range(1, 13)} | {'20'}

# Raw column of the 2024 & 2025 tab -> task path
OLDER_DATE_COLS = {
	16: 'Design>Planset sent', 17: 'Design>Planset received', 18: 'Design>Engineering stamp received',
	19: 'Zoning permit>Applied', 20: 'Zoning permit>Received',
	21: 'Building permit>Applied', 22: 'Building permit>Received',
	23: 'Interconnection>Applied (app and fee)', 25: 'Interconnection>Approved / permission to install',
	26: 'Installation>Started', 27: 'Utility>COC submitted', 28: 'Utility>PTO received',
	29: 'SREC registration', 31: 'SolarInsure registration', 32: 'SolarEdge extended warranty',
}
OLDER_PRE_PTO = [16, 17, 18, 19, 20, 21, 22, 23, 25, 26, 27]
OLDER_POST_PTO = [29, 31, 32]
OLDER_LABELS = {15: 'Deposit', 18: 'Eng. stamp', 24: 'Utility change required', 29: 'SREC setup', 30: 'Final payment',
				31: 'SolarInsure', 32: 'SolarEdge warranty', 13: 'VNM'}


def older_record_number(text: str) -> tuple[str, str] | None:
	"""'05 Sample Farm' -> ('05', 'Sample Farm'); '21 Example Barn' -> ('25021', ...); '25015 X' -> ('25015', 'X')."""
	m = re.match(r'^(\d{5}|\d{2})\s+(.+)$', text.strip())
	if not m:
		return None
	num = m.group(1)
	if len(num) == 2 and num not in OLDER_KEEP_SHORT:
		num = '250' + num
	return num, re.sub(r'\\', '', m.group(2)).strip()


def shift_year(d: dt.date, years: int) -> dt.date | None:
	try:
		return d.replace(year=d.year + years)
	except ValueError:
		return None


def street_key(street: str | None) -> str | None:
	"""'200 Mount Hope Road' and '200 Mt Hope Rd' -> '200 mount hope'."""
	if not street:
		return None
	s = re.sub(r'[.,]', ' ', street.lower())
	s = re.sub(r'\bmt\b', 'mount', s)
	s = re.sub(r'\b(road|rd|lane|ln|drive|dr|avenue|ave|street|st|boulevard|blvd|way|pike|court|ct)\b', ' ', s)
	s = re.sub(r'\s+', ' ', s).strip()
	s = re.sub(r'stone ?house', 'stonehouse', s)
	return s or None


def main_older(a) -> int:
	import sqlite3
	wb = openpyxl.load_workbook(a.xlsx, data_only=True)
	ws = wb[OLDER_TAB]

	existing_numbers, existing_names, customers_by_street = set(), set(), {}
	if a.db:
		db = sqlite3.connect(f'file:{a.db}?mode=ro', uri=True)
		existing_numbers = {r[0] for r in db.execute('SELECT project_number FROM projects')}
		existing_names = {r[0].lower() for r in db.execute('SELECT name FROM projects')}
		for oid, name, street, nproj in db.execute(
				"SELECT o.id, o.name, o.street, (SELECT COUNT(*) FROM projects p WHERE p.customer_id = o.id) FROM organizations o WHERE o.type = 'customer'"):
			k = street_key(street)
			if k:
				customers_by_street[k] = (name, nproj)

	contracts = []
	for p in a.contracts:
		path = Path(p)
		for f in sorted(path.glob('*.pdf') if path.is_dir() else [path]):
			manual = next((v for k, v in MANUAL_CONTRACTS.items() if k in f.name), None)
			if manual:
				contracts.append({'file': f.name, 'payment_option': None, 'rate_note': None, 'other_items': [], **manual,
								  'flags': list(manual['flags'])})
				continue
			try:
				contracts.append(parse_contract(f))
			except Exception as e:
				contracts.append({'file': f.name, 'error': str(e), 'flags': [f'Could not read contract: {e}']})

	projects, skipped, year_fixes = [], [], []
	for r in ws.iter_rows(min_row=4, values_only=True):
		r = list(r) + [None] * 40
		name_cell = clean(r[0])
		if not name_cell:
			continue
		parsed = older_record_number(name_cell)
		if not parsed:
			skipped.append((name_cell, 'No project number on the sheet (skipped by Greg\'s decision)'))
			continue
		number, pname = parsed
		if number in existing_numbers:
			skipped.append((name_cell, f'{number} already exists in the tracker'))
			continue
		flags, notes = [], []
		m = re.match(r'^(.*?)\s*\((\d{5,})\)$', pname)       # "Marvin Gehman (7993068)"
		if m:
			pname = m.group(1)
			notes.append(f'Sheet name included "({m.group(2)})"')
		if pname.lower() in existing_names:
			flags.append(f'Project name "{pname}" is already used by another project; importer will skip it')

		signed = iso(r[1])
		if not signed and clean(r[1]):
			flags.append(f'Contract signed cell is "{clean(r[1])}", not a date')
		sales = SALES.get(clean(r[2]).upper())
		if clean(r[2]) and not sales:
			flags.append(f'Unknown salesperson "{clean(r[2])}"')
		sheet_kw = float(r[3]) if isinstance(r[3], (int, float)) else None
		ctype, ag = CUST_TYPE.get(clean(r[4]).upper(), (None, False))
		status_raw = clean(r[14])

		# ---- dates (only cells holding a real date)
		cells: dict[int, dt.date] = {}
		other_vals: list[str] = []
		for col in list(OLDER_DATE_COLS) + [15, 24, 30, 13]:
			v = r[col]
			if v is None or clean(v) == '':
				continue
			if isinstance(v, (dt.datetime, dt.date)) and col in OLDER_DATE_COLS:
				cells[col] = v.date() if isinstance(v, dt.datetime) else v
			elif isinstance(v, (dt.datetime, dt.date)):
				other_vals.append(f'{OLDER_LABELS.get(col, col)} {iso(v)}')
			else:
				s = clean(v)
				s = s[:-2] if re.fullmatch(r'\d+\.0', s) else s
				label = OLDER_LABELS.get(col) or OLDER_DATE_COLS.get(col, str(col)).replace('>', ' > ')
				other_vals.append(f'{label}: {s}')
		for (num, path), d in OLDER_DATE_SET.items():
			if num == number:
				col = next(c for c, pth in OLDER_DATE_COLS.items() if pth == path)
				cells[col] = dt.date.fromisoformat(d)
				notes.append(f'{path.replace(">", " > ")}: {d} (from Greg)')

		# ---- two-digit-year slips: move a date by one year when that puts it in sequence
		s_date = dt.date.fromisoformat(signed) if signed else None
		pto = cells.get(28)
		if pto and s_date and pto < s_date:
			cand = shift_year(pto, 1)
			if cand and cand <= TODAY:
				year_fixes.append((number, pname, 'Utility > PTO received', pto, cand, 'before the signed date'))
				cells[28] = pto = cand
		if pto and pto > TODAY:
			cand = shift_year(pto, -1)
			if cand and (not s_date or cand >= s_date):
				year_fixes.append((number, pname, 'Utility > PTO received', pto, cand, 'in the future'))
				cells[28] = pto = cand
		lo = s_date - dt.timedelta(days=60) if s_date else None
		hi = pto or TODAY
		for col in OLDER_PRE_PTO:
			d = cells.get(col)
			if not d:
				continue
			label = OLDER_DATE_COLS[col].replace('>', ' > ')
			ok = lambda x: x is not None and (lo is None or x >= lo) and x <= hi
			if ok(d):
				continue
			why = f'after PTO {pto}' if pto and d > pto else ('in the future' if d > TODAY else f'well before the signed date {signed}')
			cand = shift_year(d, -1 if d > hi else 1)
			if ok(cand):
				year_fixes.append((number, pname, label, d, cand, why))
				cells[col] = cand
			else:
				flags.append(f'{label}: {d} is {why}; left as-is, check it')
		for col in OLDER_POST_PTO:
			d = cells.get(col)
			if not d:
				continue
			label = OLDER_DATE_COLS[col].replace('>', ' > ')
			if d > TODAY:
				cand = shift_year(d, -1)
				if cand and (not s_date or cand >= s_date):
					year_fixes.append((number, pname, label, d, cand, 'in the future'))
					cells[col] = cand
				else:
					flags.append(f'{label}: {d} is in the future; left as-is')
			elif s_date and d < s_date:
				cand = shift_year(d, 1)
				if cand and cand <= TODAY:
					year_fixes.append((number, pname, label, d, cand, f'before the signed date {signed}'))
					cells[col] = cand
				else:
					flags.append(f'{label}: {d} is before the signed date; left as-is')
		# applied/received pairs out of order after the fixes
		for a_col, b_col in ((16, 17), (19, 20), (21, 22), (23, 25), (25, 26), (26, 28), (27, 28)):
			if a_col in cells and b_col in cells and cells[a_col] > cells[b_col]:
				flags.append(f'{OLDER_DATE_COLS[a_col].replace(">", " > ")} {cells[a_col]} is after '
							 f'{OLDER_DATE_COLS[b_col].replace(">", " > ")} {cells[b_col]}')

		tasks = {OLDER_DATE_COLS[c]: {'done': d.isoformat(), 'needed': 1} for c, d in sorted(cells.items())}

		# ---- status
		hold_state = hold_reason = None
		if re.search(r'cancel', status_raw, re.I):
			hold_state, hold_reason = 'cancelled', status_raw
		completed = bool(cells.get(28)) and hold_state is None
		if not completed and hold_state is None:
			flags.append('No PTO date and not cancelled: imported as an active project (it will show open items)')

		# ---- sheet equipment, municipality, agency
		eq, eq_flags = parse_equipment(clean(r[8]), sheet_kw)
		muni, muni_flags, muni_agency = parse_ahj(clean(r[11]))
		agency, ag_flags = parse_agency(clean(r[12]))
		agency = agency or muni_agency
		util = UTILITY.get(clean(r[10]).upper().replace('  ', ' '))
		if clean(r[10]) and not util:
			flags.append(f'Unknown utility "{clean(r[10])}"')
		funding = [FUNDING[k] for k in FUNDING if k in clean(r[7]).upper()]
		tax = {'YES': 1, 'NO': 0}.get(clean(r[6]).upper())

		projects.append({
			'number': number, 'name': pname, 'customer': re.sub(r'\s*\(.*?\)\s*', ' ', pname).strip(),
			'record_only': True, 'completed': completed, 'archived': completed,
			'signed': signed, 'salesperson_email': sales, 'customer_type': ctype, 'is_agricultural': ag,
			'install_type': INSTALL_TYPE.get(clean(r[5]).upper()), 'sheet_kw': sheet_kw, 'tax_exempt': tax,
			'funding': funding, 'utility': util, 'municipality': muni, 'inspection_agency': agency,
			'status_note': None, 'hold_state': hold_state, 'hold_reason': hold_reason, 'drive_url': None,
			'has_batteries': bool(eq['batteries']), 'has_pv': True, 'equipment': eq, 'equipment_text': clean(r[8]) or None,
			'tasks': tasks, 'contact': {'name': None, 'email': clean(r[9]) or None, 'phone': None}, 'site': None,
			'price': None, 'annual_kwh': None, 'quote_number': None, 'payments': None,
			'_sheet_status': status_raw, '_other_vals': other_vals, '_eq_flags': eq_flags, '_muni_flags': muni_flags + ag_flags,
			'_ahj_text': clean(r[11]),
			'source_notes': notes, 'flags': flags,
		})

	# ---- match contracts (same greedy name matching as the main import)
	cand = []
	for ci, c in enumerate(contracts):
		if 'error' in c:
			continue
		ctoks = name_tokens(re.sub(r'\(.*?\)', '', c['file'].rsplit('.', 1)[0])) | name_tokens(c.get('contact_name') or '')
		for p in projects:
			ptoks = name_tokens(p['name'])
			sc = len(ctoks & ptoks)
			if sc:
				cand.append((sc, -len(ptoks - ctoks), ci, p['number']))
	cand.sort(reverse=True)
	used_c, used_p, by_num = set(), set(), {p['number']: p for p in projects}
	for sc, _, ci, num in cand:
		if ci in used_c or num in used_p:
			continue
		used_c.add(ci)
		used_p.add(num)
		by_num[num]['_contract'] = contracts[ci]
	unmatched = [c['file'] for i, c in enumerate(contracts) if i not in used_c]

	for p in projects:
		c = p.pop('_contract', None)
		eq = p['equipment']
		if c:
			p['contract_file'] = c['file']
			p['source_notes'].append(f"Contract: {c['file']}")
			p['flags'] += [f'Contract: {f}' for f in c.get('flags', [])
						   if not (f.startswith('Signed date read') and c.get('signed') == p['signed'])]
			p['price'] = c.get('price')
			p['annual_kwh'] = c.get('annual_kwh')
			p['quote_number'] = c.get('quote_number')
			addr = dict(c.get('address') or {})
			p['site'] = addr if addr.get('street') else None
			cn = c.get('contact_name')
			if cn and p['name'].lower().startswith(cn.lower()) and len(cn) < len(p['name']):
				cn = p['name']                                   # "Lyle" -> "Lyle Musser"
			sheet_email = p['contact']['email']
			p['contact'] = {'name': cn or p['customer'], 'email': sheet_email or c.get('email'), 'phone': c.get('phone')}
			if sheet_email and c.get('email') and sheet_email.lower() != c['email'].lower():
				p['flags'].append(f"Email: sheet {sheet_email}, contract {c['email']} (kept the sheet's)")
			if c.get('signed') and not p['signed']:
				p['signed'] = c['signed']
				p['source_notes'].append('Signed date from the contract (sheet blank)')
			elif c.get('signed') and c['signed'] != p['signed']:
				p['flags'].append(f"Signed date: sheet {p['signed']}, contract {c['signed']} (kept the sheet's)")
			sp = c.get('salesperson')
			if sp and not p['salesperson_email']:
				p['salesperson_email'] = SALES.get(sp.split()[0].upper())
			elif sp and p['salesperson_email'] and sp.split()[0].lower() not in p['salesperson_email']:
				p['flags'].append(f"Salesperson: sheet {p['salesperson_email']}, contract {sp} (kept the sheet's)")
			for f in c.get('funding', []):
				if f not in p['funding']:
					p['funding'].append(f)
					p['source_notes'].append(f'Funding from contract: {f}')
			if c.get('rate_note'):
				p['source_notes'].append(f"Contract savings basis: {c['rate_note']}")
			if c.get('payment_option'):
				p['source_notes'].append(f"Payment option: {c['payment_option']}")
			for item in c.get('other_items', []):
				p['source_notes'].append(f'Contract also includes: {item}')
			# Equipment: the contract's lines (full model names) when its DC total matches the sheet kW
			c_dc = sum(m['qty'] * m['watts'] for m in c.get('modules', [])) / 1000
			fmt_inv = lambda lines: '; '.join(f"{i['qty']}x{i['ac_kw']:g}kW" for i in lines) or 'none'
			fmt_mod = lambda lines: '; '.join(f"{m['qty']}x{m['watts']:g}W" for m in lines) or 'none'
			if c_dc and p['sheet_kw'] and abs(c_dc - p['sheet_kw']) / p['sheet_kw'] <= 0.01:
				if fmt_mod(eq['modules']) != fmt_mod(c['modules']):
					p['flags'].append(f"Modules: sheet text {fmt_mod(eq['modules'])}, contract {fmt_mod(c['modules'])} "
									  f"(used the contract; it matches the sheet's {p['sheet_kw']:g} kW)")
				if c.get('inverters') and sorted(fmt_inv(eq['inverters']).split('; ')) != sorted(fmt_inv(c['inverters']).split('; ')):
					p['flags'].append(f"Inverters: sheet text {fmt_inv(eq['inverters'])}, contract {fmt_inv(c['inverters'])} (used the contract)")
				eq['modules'] = c['modules']
				if c.get('inverters'):
					eq['inverters'] = c['inverters']
				if c.get('batteries'):
					eq['batteries'] = c['batteries']
				if c.get('racking'):
					eq['racking'] = '; '.join(c['racking'])
				p['_eq_flags'] = []
			elif c_dc:
				p['flags'].append(f"Contract equipment ({fmt_mod(c['modules'])}, {fmt_inv(c.get('inverters', []))}, "
								  f"{c_dc:.2f} kW) does not match the sheet's {p['sheet_kw']} kW; kept the sheet's equipment")
			p['has_batteries'] = bool(eq['batteries'])
		elif p['hold_state'] != 'cancelled':
			p['flags'].append('No contract: price, site address and contact not filled')
		p['flags'] += p.pop('_eq_flags')
		n = p['number']
		if n in OLDER_SITE and not p.get('site'):
			p['site'] = OLDER_SITE[n]
			p['source_notes'].append('Site address from Greg')
			p['flags'] = [f for f in p['flags'] if not f.startswith('No contract')]
			p['flags'].append('No contract: price and contact not filled')
		if n in OLDER_EQUIPMENT:
			eq.update(OLDER_EQUIPMENT[n])
		for (num, prefix), note in OLDER_ANSWERED.items():
			if num == n and any(f.startswith(prefix) for f in p['flags']):
				p['flags'] = [f for f in p['flags'] if not f.startswith(prefix)]
				if note:
					p['source_notes'].append(note)

		# Municipality: Census for the contract address, else the sheet's AHJ
		n = p['number']
		sheet_muni = p['municipality']
		if n in OLDER_CENSUS:
			name, county, zipc = OLDER_CENSUS[n]
			p['municipality'] = {'name': name, 'county': county, 'state': 'PA'}
			if p['site'] and not p['site'].get('zip'):
				p['site'].update({'state': 'PA', 'zip': zipc})
				p['source_notes'].append('Site ZIP from the Census address match')
			if sheet_muni and (sheet_muni['name'].lower(), sheet_muni['county'].lower()) != (name.lower(), county.lower()):
				p['flags'].append(f"Municipality: sheet AHJ \"{p['_ahj_text']}\"; Census {name}, {county} Co. (used Census)")
			else:
				p['source_notes'].append('Municipality from Census address lookup')
			if n in OLDER_MUNI_NOTES:
				p['flags'].append(OLDER_MUNI_NOTES[n])
			p.pop('_muni_flags')
		elif n in OLDER_MUNI_CONFIRMED:
			name, county = OLDER_MUNI_CONFIRMED[n]
			p['municipality'] = {'name': name, 'county': county, 'state': 'PA'}
			p['source_notes'].append('Municipality confirmed by Greg')
			p.pop('_muni_flags')
		else:
			p['flags'] += p.pop('_muni_flags')
			if not p['municipality'] and p['hold_state'] != 'cancelled':
				p['flags'].append('Municipality unknown (no contract address and no AHJ on the sheet)')

		# Customer: link to an existing customer at the same street address (service records)
		k = street_key((p.get('site') or {}).get('street'))
		if k and k in customers_by_street:
			cname, nproj = customers_by_street[k]
			if cname != p['customer']:
				p['source_notes'].append(f'Customer is the existing record "{cname}" (same street address)')
				p['customer'] = cname

		# Import note: what the sheet said that did not become a dated step
		if p['_sheet_status']:
			p['source_notes'].append(f"Sheet status: {p['_sheet_status']}")
		if p['_other_vals']:
			p['source_notes'].append('Sheet values without a date (not imported as steps): ' + '; '.join(p['_other_vals']))
		p.pop('_sheet_status')
		p.pop('_ahj_text')
		p.pop('_other_vals')
		p['source_notes'].insert(0, f'Record-only import from the {OLDER_TAB} tab: only the dates on the sheet became steps; '
									'no payments or blank steps were created'
									+ ('; PTO date present, so imported as Completed' if p['completed'] else ''))

	out = {'generated': dt.datetime.now().isoformat(timespec='seconds'), 'source': Path(a.xlsx).name,
		   'mode': 'record_only', 'projects': projects}
	Path(a.out).parent.mkdir(parents=True, exist_ok=True)
	Path(a.out).write_text(json.dumps(out, indent=1, default=str))

	# ---- report
	money = lambda v: f'${v:,.2f}' if v is not None else '(none)'
	L = [f'# Record-only import review: {len(projects)} projects from the "{OLDER_TAB}" tab', '',
		 f"Generated {out['generated']}. {len(contracts)} contract(s); existing projects checked against "
		 f"{Path(a.db).name if a.db else '(no database given)'}.", '',
		 '| # | Project | Signed | PTO | Status | Price | kW DC | Municipality | Contract |', '|---|---|---|---|---|---|---|---|---|']
	for p in projects:
		dc = sum(m['qty'] * m['watts'] for m in p['equipment']['modules']) / 1000
		mu = f"{p['municipality']['name']}, {p['municipality']['county']} Co." if p['municipality'] else ''
		st = 'Cancelled' if p['hold_state'] == 'cancelled' else 'Completed' if p['completed'] else 'Active'
		L.append(f"| {p['number']} | {p['name']} | {p['signed'] or ''} | {p['tasks'].get('Utility>PTO received', {}).get('done', '')} "
				 f"| {st} | {money(p['price']) if p['price'] else ''} | {dc:.2f} | {mu} | {'yes' if p.get('contract_file') else 'no'} |")
	L.append('')
	if year_fixes:
		L += ['## Years corrected (one-year shift put the date back in sequence)', '',
			  '| # | Project | Step | Sheet | Imported | Why |', '|---|---|---|---|---|---|']
		L += [f'| {n} | {nm} | {lab} | {a_} | {b_} | {why} |' for n, nm, lab, a_, b_, why in year_fixes] + ['']
	if skipped:
		L += ['## Not imported', ''] + [f'- **{n}**: {why}' for n, why in skipped] + ['']
	if unmatched:
		L += ['## Contracts that matched no row', ''] + [f'- {f}' for f in unmatched] + ['']
	L += ['## Per project review notes', '']
	for p in projects:
		if not p['flags']:
			continue
		L.append(f"### {p['number']} {p['name']}")
		L += [f'- {f}' for f in p['flags']] + ['']
	Path(a.report).write_text('\n'.join(L))
	print(f"{len(projects)} projects ({sum(p['completed'] for p in projects)} completed, "
		  f"{sum(p['hold_state'] == 'cancelled' for p in projects)} cancelled), {len(skipped)} skipped, "
		  f"{len(year_fixes)} year fixes, {sum(1 for p in projects if p['flags'])} with review notes -> {a.out}, {a.report}")
	return 0


if __name__ == '__main__':
	if '--older' in sys.argv:
		ap = argparse.ArgumentParser(description=f'Record-only import of the "{OLDER_TAB}" tab')
		ap.add_argument('--older', action='store_true')
		ap.add_argument('--xlsx', required=True)
		ap.add_argument('--contracts', nargs='*', default=[])
		ap.add_argument('--db', help='a tracker backup (read-only) to skip existing projects and match customers')
		ap.add_argument('--out', required=True)
		ap.add_argument('--report', required=True)
		sys.exit(main_older(ap.parse_args()))
	sys.exit(main())
